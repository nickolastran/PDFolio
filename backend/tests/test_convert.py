"""Every conversion path, plus compress/repair/OCR and the SSRF / sandbox defences of converters."""
import base64
import http.server
import io
import json
import threading
import zipfile

import docx
import openpyxl
import pikepdf
import pptx
import pymupdf
import pytest
from PIL import Image

from app.ops import convert
from conftest import make_pdf, run_job


def _docx(text="Hello DOCX") -> bytes:
    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = "Hello PPTX"
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Hello XLSX"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _png(size=(40, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, "red").save(buf, "PNG")
    return buf.getvalue()


def _site_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("index.html", '<h1>Hello ZIP</h1><img src="img/dot.png" style="width:100px">'
                                  '<link rel="stylesheet" href="style.css">')
        zf.writestr("style.css", "h1 { color: black }")
        zf.writestr("img/dot.png", _png())
    return buf.getvalue()


def _table_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for r in range(3):
        for c in range(2):
            rect = pymupdf.Rect(72 + c * 100, 72 + r * 30, 172 + c * 100, 102 + r * 30)
            page.draw_rect(rect, color=(0, 0, 0), width=1)
            page.insert_text((rect.x0 + 5, rect.y0 + 20), "=1+1" if (r, c) == (2, 1) else f"R{r}C{c}")
    return doc.tobytes()


SAMPLES = {
    "docx": ("a.docx", _docx(), "Hello DOCX"),
    "rtf": ("a.rtf", rb"{\rtf1\ansi Hello RTF}", "Hello RTF"),
    "pptx": ("a.pptx", _pptx(), "Hello PPTX"),
    "xlsx": ("a.xlsx", _xlsx(), "Hello XLSX"),
    "md": ("a.md", b"# Hello MD\n\nSome *text*.", "Hello MD"),
    "html": ("a.html", b"<!doctype html><h1>Hello HTML</h1>", "Hello HTML"),
    "txt": ("a.txt", b"Hello TXT\nline two", "Hello TXT"),
    "htmlzip": ("site.zip", _site_zip(), "Hello ZIP"),
    "pdf": ("a.pdf", make_pdf("Hello PDF"), "Hello PDF"),
}


def convert_via_api(client, kind, to, tool="convert", **opts):
    name, data, _ = SAMPLES[kind]
    job = run_job(client, tool, [(name, data)], json.dumps({"to": to, **opts} if tool == "convert" else opts))
    r = client.get(f"/api/jobs/{job}/download")
    return r


def pdf_text(data: bytes) -> str:
    return " ".join(p.get_text() for p in pymupdf.open(stream=data, filetype="pdf"))


@pytest.mark.parametrize("kind", ["docx", "rtf", "pptx", "xlsx", "md", "html", "txt", "htmlzip"])
def test_to_pdf(client, kind):
    r = convert_via_api(client, kind, "pdf")
    assert r.headers["content-type"] == "application/pdf"
    assert SAMPLES[kind][2] in pdf_text(r.content)


def test_htmlzip_includes_images_from_the_zip(client):
    r = convert_via_api(client, "htmlzip", "pdf")
    assert pymupdf.open(stream=r.content, filetype="pdf")[0].get_images()


@pytest.mark.parametrize("src,dst", [
    ("md", "docx"), ("md", "html"), ("md", "odt"), ("md", "rtf"), ("md", "epub"), ("md", "txt"),
    ("html", "md"), ("docx", "md"), ("docx", "odt"), ("rtf", "docx"), ("txt", "md"), ("htmlzip", "html"),
])
def test_document_conversions(client, src, dst):
    r = convert_via_api(client, src, dst)
    assert r.headers["content-disposition"].endswith(f'.{dst}"')
    text = SAMPLES[src][2]
    data = r.content
    if dst in ("docx", "odt", "epub"):
        data = b"".join(zipfile.ZipFile(io.BytesIO(data)).read(n) for n in zipfile.ZipFile(io.BytesIO(data)).namelist())
    assert text.encode() in data


def test_epub_round_trip(client):
    epub = convert_via_api(client, "md", "epub").content
    job = run_job(client, "convert", [("book.epub", epub)], '{"to": "pdf"}')
    assert "Hello MD" in pdf_text(client.get(f"/api/jobs/{job}/download").content)


def test_odt_to_pdf(client):
    odt = convert_via_api(client, "docx", "odt").content
    job = run_job(client, "word-to-pdf", [("a.odt", odt)])
    r = client.get(f"/api/jobs/{job}/download")
    assert r.headers["content-disposition"] == 'attachment; filename="a.pdf"'
    assert "Hello DOCX" in pdf_text(r.content)


def test_pdf_to_word(client):
    r = convert_via_api(client, "pdf", None, tool="pdf-to-word")
    assert "Hello PDF" in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)


def test_pdf_to_pptx(client):
    r = convert_via_api(client, "pdf", None, tool="pdf-to-pptx")
    assert len(pptx.Presentation(io.BytesIO(r.content)).slides) == 1


def test_pdf_to_xlsx_tables_and_no_formulas(client):
    job = run_job(client, "pdf-to-xlsx", [("t.pdf", _table_pdf())])
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/jobs/{job}/download").content))
    cells = [c for row in wb.worksheets[0].iter_rows() for c in row if c.value]
    assert {c.value for c in cells} >= {"R0C0", "=1+1"}
    assert all(c.data_type == "s" for c in cells)   # "=1+1" is stored as text, not a formula


def test_pdf_to_xlsx_without_tables(client):
    job = run_job(client, "pdf-to-xlsx", [("a.pdf", make_pdf("no table"))], expect="failed")
    assert client.get(f"/api/jobs/{job}").json()["error"] == "No tables found in this PDF"


def test_pdf_to_markdown_and_txt(client):
    assert "Hello PDF" in convert_via_api(client, "pdf", None, tool="pdf-to-md").text
    assert "Hello PDF" in convert_via_api(client, "pdf", "txt").text


def test_pdf_to_pdfa(client):
    r = convert_via_api(client, "pdf", None, tool="pdf-to-pdfa")
    assert "Hello PDF" in pdf_text(r.content)   # Ghostscript keeps the text extractable
    with pikepdf.open(io.BytesIO(r.content)) as pdf:
        assert pdf.open_metadata().pdfa_status.startswith("2")


def test_pdf_to_jpg_single_and_multi(client):
    r = convert_via_api(client, "pdf", None, tool="pdf-to-jpg", dpi=72)
    assert r.headers["content-type"] == "image/jpeg"
    assert Image.open(io.BytesIO(r.content)).size == (595, 842)
    job = run_job(client, "pdf-to-jpg", [("a.pdf", make_pdf("1", "2", "3"))])
    r = client.get(f"/api/jobs/{job}/download")
    assert r.headers["content-type"] == "application/zip"
    assert zipfile.ZipFile(io.BytesIO(r.content)).namelist() == ["page-001.jpg", "page-002.jpg", "page-003.jpg"]


def test_render_dpi_is_capped(client):
    r = convert_via_api(client, "pdf", "png", dpi=600)   # MAX_RENDER_DPI is 300
    w, h = Image.open(io.BytesIO(r.content)).size
    assert abs(w - 595 * 300 / 72) <= 1 and abs(h - 842 * 300 / 72) <= 1


def test_images(client):
    job = run_job(client, "convert", [("a.png", _png())], '{"to": "jpg"}')
    assert client.get(f"/api/jobs/{job}/download").headers["content-type"] == "image/jpeg"
    job = run_job(client, "convert", [("a.png", _png())], '{"to": "pdf"}')
    assert pymupdf.open(stream=client.get(f"/api/jobs/{job}/download").content, filetype="pdf")[0].get_images()


def test_impossible_conversions_are_rejected(client):
    f = [("files", ("a.md", b"# x"))]
    assert client.post("/api/jobs", data={"tool": "convert", "options": '{"to": "xlsx"}'}, files=f).status_code == 415
    assert client.post("/api/jobs", data={"tool": "convert", "options": '{"to": "htmlzip"}'}, files=f).status_code == 422
    assert client.get("/api/formats").json()["md"].count("xlsx") == 0


def test_fixed_tool_rejects_other_inputs(client):
    r = client.post("/api/jobs", data={"tool": "word-to-pdf"}, files=[("files", ("a.pdf", make_pdf("x")))])
    assert r.status_code == 415


# --- optimize ---------------------------------------------------------------

def _image_pdf(text="") -> bytes:
    """A PDF whose page is one big photo-like image, optionally of rendered text (for OCR)."""
    if text:
        src = pymupdf.open()
        src.new_page().insert_text((72, 120), text, fontsize=40)
        png = src[0].get_pixmap(dpi=200).tobytes("png")
    else:
        buf = io.BytesIO()
        Image.effect_noise((1500, 1500), 64).convert("RGB").save(buf, "PNG")
        png = buf.getvalue()
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=png)
    return doc.tobytes()


def test_compress_makes_it_smaller(client):
    data = _image_pdf()
    job = run_job(client, "compress", [("a.pdf", data)], '{"level": "high"}')
    assert len(client.get(f"/api/jobs/{job}/download").content) < len(data) / 2


def test_repair_broken_pdf(client):
    broken = make_pdf("Fix me")
    broken = broken[: broken.rindex(b"xref")] + b"%%EOF"   # chop the xref table and trailer
    job = run_job(client, "repair", [("a.pdf", broken)])
    assert "Fix me" in pdf_text(client.get(f"/api/jobs/{job}/download").content)


def test_ocr_makes_text_searchable(client):
    data = _image_pdf("Searchable words")
    assert pdf_text(data).strip() == ""
    job = run_job(client, "ocr", [("scan.pdf", data)])
    assert "Searchable" in pdf_text(client.get(f"/api/jobs/{job}/download").content)


# --- converter security -------------------------------------------------------

@pytest.mark.parametrize("url", [
    "http://169.254.169.254/latest/meta-data/", "http://localhost/", "http://10.0.0.1/x.png",
    "https://example.com/x.png", "file:///etc/passwd", "ftp://example.com/x",
])
def test_fetcher_blocks_external(url, tmp_path):
    with pytest.raises(ValueError):
        convert.SafeFetcher(tmp_path).fetch(url)


def test_fetcher_blocks_escape_from_zip_root(tmp_path):
    (tmp_path / "site").mkdir()
    (tmp_path / "secret.txt").write_text("x")
    (tmp_path / "site" / "ok.css").write_text("x")
    f = convert.SafeFetcher(tmp_path / "site")
    f.fetch((tmp_path / "site" / "ok.css").as_uri())
    with pytest.raises(ValueError):
        f.fetch((tmp_path / "site" / ".." / "secret.txt").as_uri())


def test_html_to_pdf_makes_no_requests(tmp_path):
    """Outside the sandbox (so the network is reachable), WeasyPrint still fetches nothing."""
    hits = []

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(200)
            self.end_headers()

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    port = srv.server_address[1]
    page = tmp_path / "in.html"
    page.write_text(f'<img src="http://127.0.0.1:{port}/img.png"><link rel=stylesheet href="http://127.0.0.1:{port}/s.css">'
                    '<img src="http://169.254.169.254/latest/meta-data/"><img src="file:///etc/passwd">'
                    f'<style>@import url("http://127.0.0.1:{port}/i.css"); body{{background:url(file:///etc/hostname)}}</style>Hi')
    out = convert.weasyprint(page, "html", "pdf", tmp_path, None)
    srv.shutdown()
    assert hits == []
    assert "Hi" in pdf_text(out.read_bytes())


def test_pandoc_sandbox_cannot_read_local_files(tmp_path):
    secret = tmp_path / "secret.png"
    secret.write_bytes(_png())
    md = tmp_path / "in.md"
    md.write_text(f"![x]({secret})\n\n```{{.include}}\n{secret}\n```\n")
    out = convert.pandoc(md, "md", "html", tmp_path, None)
    assert base64.b64encode(secret.read_bytes())[:40] not in out.read_bytes()


def test_xxe_docx_via_api(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_docx())) as src, zipfile.ZipFile(buf, "w") as dst:
        for n in src.namelist():
            data = src.read(n)
            if n == "word/document.xml":
                data = data.replace(b"?>", b'?><!DOCTYPE w [<!ENTITY x SYSTEM "file:///etc/passwd">]>', 1)
            dst.writestr(n, data)
    job = run_job(client, "word-to-pdf", [("a.docx", buf.getvalue())], expect="failed")
    assert client.get(f"/api/jobs/{job}").json()["error"] == "File contains unsafe XML"
